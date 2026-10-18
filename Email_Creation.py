import os
import re
import pandas as pd
import streamlit as st
import tkinter as tk
from tkinter import filedialog
import pythoncom
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication

# Import docx for Word file body parsing
try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# Try importing pywin32 for Outlook automation on Windows
try:
    import win32com.client as win32
    OUTLOOK_AVAILABLE = True
except ImportError:
    OUTLOOK_AVAILABLE = False


def select_folder():
    """Opens a native desktop folder picker dialog."""
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    folder_path = filedialog.askdirectory(master=root)
    root.destroy()
    return folder_path


def build_html_from_body_file(uploaded_file, context):
    """Converts uploaded Excel or Word body file into HTML and replaces recipient placeholders."""
    filename = uploaded_file.name.lower()
    
    # 1. Handle Excel Body File (.xlsx, .xls)
    if filename.endswith(('.xlsx', '.xls')):
        try:
            uploaded_file.seek(0)
            df_body = pd.read_excel(uploaded_file, header=None)
            
            html_table = [
                '<table border="1" cellpadding="8" cellspacing="0" '
                'style="border-collapse: collapse; border-color: #CBD5E1; font-family: Arial, sans-serif; font-size: 13px; width: 100%; max-width: 700px;">'
            ]
            
            for _, row in df_body.iterrows():
                val_a = str(row[0]).strip() if pd.notna(row[0]) else ""
                val_b = str(row[1]).strip() if len(row) > 1 and pd.notna(row[1]) else ""
                
                # Dynamic placeholder replacement
                for k, v in context.items():
                    val_a = val_a.replace(f"{{{k}}}", v)
                    val_b = val_b.replace(f"{{{k}}}", v)
                
                val_a_html = val_a.replace('\n', '<br>')
                val_b_html = val_b.replace('\n', '<br>')
                
                if val_a and not val_b:
                    html_table.append(
                        f'<tr style="background-color: #E2E8F0;"><td colspan="2" style="padding: 8px; font-weight: bold; color: #0F172A;">{val_a_html}</td></tr>'
                    )
                else:
                    html_table.append(
                        f'<tr><td style="padding: 8px; font-weight: bold; width: 40%; background-color: #F8FAFC; color: #334155;">{val_a_html}</td>'
                        f'<td style="padding: 8px; width: 60%; color: #0F172A;">{val_b_html}</td></tr>'
                    )
            
            html_table.append('</table>')
            
            insured_name = context.get("Insured Name", "")
            full_html = (
                f"<p>Dear <b>{insured_name}</b>,</p>"
                "<p>Please find attached your motor insurance policy details and terms below:</p>"
                + "".join(html_table) +
                "<p style='margin-top:15px;'>Best regards,<br><b>Operations Team</b></p>"
            )
            return full_html
        except Exception as e:
            st.error(f"Error parsing Excel body file: {e}")
            return None

    # 2. Handle Word Document Body File (.docx)
    elif filename.endswith('.docx'):
        if not DOCX_AVAILABLE:
            st.error("python-docx package is missing. Run `pip install python-docx` to enable .docx body files.")
            return None
        try:
            uploaded_file.seek(0)
            doc = docx.Document(uploaded_file)
            body_paragraphs = []
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    for k, v in context.items():
                        p_text = p_text.replace(f"{{{k}}}", v)
                    body_paragraphs.append(f"<p>{p_text}</p>")
            return "".join(body_paragraphs)
        except Exception as e:
            st.error(f"Error parsing Word body file: {e}")
            return None

    return None


def create_eml_file(to_email, cc_list, subject, body_html, attachment_path, output_dir, file_prefix):
    """Creates a standalone .eml email file that can be opened directly in Outlook."""
    try:
        msg = MIMEMultipart()
        msg['To'] = str(to_email).strip() if pd.notna(to_email) else ""
        
        valid_cc = [str(cc).strip() for cc in cc_list if pd.notna(cc) and str(cc).strip() and str(cc).strip().lower() != 'nan']
        if valid_cc:
            msg['Cc'] = ", ".join(valid_cc)
            
        msg['Subject'] = str(subject)
        msg.attach(MIMEText(body_html, 'html'))
        
        if attachment_path and os.path.exists(attachment_path):
            filename = os.path.basename(attachment_path)
            with open(attachment_path, "rb") as f:
                part = MIMEApplication(f.read(), Name=filename)
            part['Content-Disposition'] = f'attachment; filename="{filename}"'
            msg.attach(part)
            
        # Clean file name for disk storage
        safe_prefix = re.sub(r'[^\w\s-]', '', str(file_prefix)).strip().replace(' ', '_')
        eml_filename = f"Draft_{safe_prefix}.eml"
        eml_filepath = os.path.join(output_dir, eml_filename)
        
        with open(eml_filepath, "wb") as f:
            f.write(msg.as_bytes())
            
        return True, eml_filepath
    except Exception as e:
        return False, str(e)


def create_outlook_draft(to_email, cc_list, subject, body_html, attachment_path=None):
    """Creates a draft email directly inside local Microsoft Outlook application via COM."""
    try:
        pythoncom.CoInitialize()
        
        outlook = win32.Dispatch('Outlook.Application')
        mail = outlook.CreateItem(0)  # 0 = olMailItem
        
        mail.To = str(to_email).strip() if pd.notna(to_email) else ""
        
        valid_cc = [str(cc).strip() for cc in cc_list if pd.notna(cc) and str(cc).strip() and str(cc).strip().lower() != 'nan']
        if valid_cc:
            mail.CC = "; ".join(valid_cc)
            
        mail.Subject = str(subject)
        mail.HTMLBody = str(body_html)
        
        if attachment_path and os.path.exists(attachment_path):
            abs_attachment_path = os.path.abspath(attachment_path)
            mail.Attachments.Add(abs_attachment_path)
            
        mail.Save()  # Saves as draft in Outlook
        return True, "Draft saved in Outlook"
    except Exception as e:
        return False, str(e)
    finally:
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


# --- Streamlit UI Configuration ---
st.set_page_config(
    page_title="Motor Insurance Policy Email Automation",
    page_icon="🚗",
    layout="wide"
)

# Custom Styling
st.markdown("""
    <style>
        .main-header {
            font-size: 1.8rem;
            font-weight: 700;
            color: #1E293B;
            margin-bottom: 0.2rem;
        }
        .sub-header {
            font-size: 0.95rem;
            color: #64748B;
            margin-bottom: 1.5rem;
        }
        .stButton>button {
            width: 100%;
            background-color: #2563EB;
            color: white;
            font-weight: 600;
            border-radius: 8px;
            padding: 0.6rem 1rem;
        }
    </style>
""", unsafe_allow_html=True)


# Session State Initializations
if "pdf_folder_path" not in st.session_state:
    st.session_state.pdf_folder_path = ""
if "eml_folder_path" not in st.session_state:
    st.session_state.eml_folder_path = ""


# ==========================================
# 1. LEFT PANEL: SETTINGS & MAPPING GUIDELINES
# ==========================================
with st.sidebar:
    st.subheader("⚙️ Configuration & Guidelines")
    st.markdown("---")
    
    st.markdown("### 📊 Master Data Columns")
    st.info(
        "• **Col A:** Insured Name *(PDF Key)*\n"
        "• **Col B:** Mobile Number\n"
        "• **Col C:** Insurance Company\n"
        "• **Col D:** Policy No.\n"
        "• **Col E:** Vehicle Type / Year\n"
        "• **Col F:** Chassis No.\n"
        "• **Col G:** Vehicle Value\n"
        "• **Col H:** Cover Date\n"
        "• **Col I:** Email Address *(To)*\n"
        "• **Col J to L:** cc1, cc2, cc3"
    )
    
    st.markdown("### 📄 Email Body File Upload")
    st.success(
        "Upload an Excel or Word file containing your terms and requirements table. Placeholders like {Insured Name} and {Policy No.} are mapped automatically."
    )
    
    st.markdown("---")
    st.caption("Environment Status:")
    if OUTLOOK_AVAILABLE:
        st.success("🟢 Outlook COM Active (with .EML fallback)")
    else:
        st.info("🔵 File Mode (.EML file creation)")


# ==========================================
# 2. MAIN PANEL: STAGE 1 & STAGE 2
# ==========================================
st.markdown('<div class="main-header">🚗 Motor Insurance Quote & Policy Email Automation</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-header">Automate personalized email drafts, attach policy PDFs, and generate openable draft files.</div>', unsafe_allow_html=True)

# --- STAGE 1: INPUT FORM & SETUP ---
st.markdown("### 🛠️ Stage 1: Template & File Configurations")

with st.form("draft_form"):
    email_subject = st.text_input(
        "Email Subject Line:",
        value="Motor Insurance Quotation & Policy - {Policy No.} - {Insured Name}"
    )
    
    st.markdown("#### 📝 Choose Email Body Source")
    body_source_option = st.radio(
        "Select Email Body Option:",
        options=["Upload Custom Email Body File (.xlsx, .docx)", "Use Default Editor Template"],
        horizontal=True
    )
    
    col_b1, col_b2 = st.columns(2)
    
    with col_b1:
        uploaded_body_file = st.file_uploader(
            "Upload Email Body Content File (.xlsx or .docx):",
            type=["xlsx", "xls", "docx"],
            help="Upload your Excel sheet containing the terms/table or a Word doc."
        )
        
    with col_b2:
        default_editor_body = st.text_area(
            "Fallback / Editor Body Template:",
            value="Dear {Insured Name},\n\nPlease find attached your motor insurance quotation for Policy No: {Policy No.}.\n\nBest regards,\nOperations Team",
            height=120
        )
    
    st.divider()
    
    col_u1, col_u2 = st.columns(2)
    
    with col_u1:
        master_excel_file = st.file_uploader(
            "Upload Master Recipient Data Excel (.xlsx, .xls)",
            type=["xlsx", "xls"],
            help="Must contain Insured Name in Column A, Policy No in Column D, Email in Column I, and CC in columns J, K, L."
        )
        
    with col_u2:
        st.write("**Select PDF Storage Folder:**")
        sub_c1, sub_c2 = st.columns([3, 1])
        with sub_c1:
            st.text_input(
                "PDF Path",
                value=st.session_state.pdf_folder_path,
                placeholder="No PDF folder selected",
                disabled=True,
                label_visibility="collapsed"
            )
        with sub_c2:
            if st.form_submit_button("Browse PDFs"):
                selected = select_folder()
                if selected:
                    st.session_state.pdf_folder_path = selected
                    st.rerun()

        st.write("**Select Draft (.eml) Output Folder:**")
        sub_e1, sub_e2 = st.columns([3, 1])
        with sub_e1:
            st.text_input(
                "EML Path",
                value=st.session_state.eml_folder_path,
                placeholder="No output folder selected",
                disabled=True,
                label_visibility="collapsed"
            )
        with sub_e2:
            if st.form_submit_button("Browse Output"):
                selected = select_folder()
                if selected:
                    st.session_state.eml_folder_path = selected
                    st.rerun()

    st.divider()
    submit_btn = st.form_submit_button("🚀 Start Processing & Generate Drafts", use_container_width=True)


# --- STAGE 2: PROCESSING STATUS & TOTALS ---
if submit_btn:
    st.markdown("---")
    st.markdown("### 📈 Stage 2: Processing Status & Results")
    
    if not master_excel_file:
        st.error("⚠️ Please upload the Master Recipient Data Excel sheet to proceed.")
        st.stop()
        
    if not st.session_state.pdf_folder_path or not os.path.exists(st.session_state.pdf_folder_path):
        st.error("⚠️ Please select a valid folder containing the PDF files.")
        st.stop()

    # Default output directory to PDF folder if EML output folder is not specified
    eml_output_dir = st.session_state.eml_folder_path if st.session_state.eml_folder_path else st.session_state.pdf_folder_path

    try:
        master_excel_file.seek(0)
        df = pd.read_excel(master_excel_file)
    except Exception as e:
        st.error(f"Failed to read Master Excel file: {e}")
        st.stop()

    if len(df.columns) < 9:
        st.error(f"⚠️ Your Excel sheet has {len(df.columns)} columns, but at least 9 columns are required.")
        st.stop()

    # Map exact column positions
    c_insured_name = df.columns[0]  # Col A
    c_mobile       = df.columns[1]  # Col B
    c_insurance_co = df.columns[2]  # Col C
    c_policy_no    = df.columns[3]  # Col D
    c_vehicle_type = df.columns[4]  # Col E
    c_chassis      = df.columns[5]  # Col F
    c_vehicle_val  = df.columns[6]  # Col G
    c_cover_date   = df.columns[7]  # Col H
    c_email        = df.columns[8]  # Col I
    
    c_cc1 = df.columns[9] if len(df.columns) > 9 else None
    c_cc2 = df.columns[10] if len(df.columns) > 10 else None
    c_cc3 = df.columns[11] if len(df.columns) > 11 else None

    pdf_files = [f for f in os.listdir(st.session_state.pdf_folder_path) if f.lower().endswith(".pdf")]
    
    if not pdf_files:
        st.warning("⚠️ No PDF files found in the selected folder.")

    success_count = 0
    skipped_count = 0
    failed_count = 0
    logs = []

    progress_bar = st.progress(0)
    total_rows = len(df)

    for idx, row in df.iterrows():
        insured_name = str(row[c_insured_name]).strip() if pd.notna(row[c_insured_name]) else ""
        policy_no = str(row[c_policy_no]).strip() if pd.notna(row[c_policy_no]) else ""
        
        if not insured_name:
            logs.append({"Row": idx + 1, "Insured Name": "N/A", "Status": "Skipped", "Details": "Insured Name missing"})
            skipped_count += 1
            continue

        # Match PDF file using Insured Name (Col A)
        matched_pdf = None
        for pdf in pdf_files:
            if re.search(re.escape(insured_name), pdf, re.IGNORECASE):
                matched_pdf = os.path.join(st.session_state.pdf_folder_path, pdf)
                break

        if not matched_pdf:
            logs.append({"Row": idx + 1, "Insured Name": insured_name, "Status": "Skipped", "Details": "Matching PDF not found in folder"})
            skipped_count += 1
            continue

        to_email = row[c_email] if pd.notna(row[c_email]) else ""
        cc_list = []
        if c_cc1 and pd.notna(row[c_cc1]): cc_list.append(row[c_cc1])
        if c_cc2 and pd.notna(row[c_cc2]): cc_list.append(row[c_cc2])
        if c_cc3 and pd.notna(row[c_cc3]): cc_list.append(row[c_cc3])

        context = {
            "Insured Name": insured_name,
            "Mobile Number": str(row[c_mobile]) if pd.notna(row[c_mobile]) else "",
            "Insurance Company": str(row[c_insurance_co]) if pd.notna(row[c_insurance_co]) else "",
            "Policy No.": policy_no,
            "Vehicle Type / Year of Manufacture": str(row[c_vehicle_type]) if pd.notna(row[c_vehicle_type]) else "",
            "Chassis No.": str(row[c_chassis]) if pd.notna(row[c_chassis]) else "",
            "Vehicle Value": str(row[c_vehicle_val]) if pd.notna(row[c_vehicle_val]) else "",
            "Cover Date": str(row[c_cover_date]) if pd.notna(row[c_cover_date]) else "",
            "Email Address": str(to_email)
        }

        # Subject Line Formatting
        formatted_subject = email_subject
        for k, v in context.items():
            formatted_subject = formatted_subject.replace(f"{{{k}}}", v)

        # Body Formatting
        if body_source_option.startswith("Upload") and uploaded_body_file is not None:
            formatted_body = build_html_from_body_file(uploaded_body_file, context)
        else:
            paragraphs = default_editor_body.split('\n')
            formatted_body = "".join([f"<p>{p.strip()}</p>" for p in paragraphs if p.strip()])
            for k, v in context.items():
                formatted_body = formatted_body.replace(f"{{{k}}}", v)

        draft_created = False
        
        # 1. Attempt Outlook COM automation if available
        if OUTLOOK_AVAILABLE:
            ok, msg = create_outlook_draft(to_email, cc_list, formatted_subject, formatted_body, matched_pdf)
            if ok:
                success_count += 1
                draft_created = True
                logs.append({"Row": idx + 1, "Insured Name": insured_name, "Status": "Success (Outlook)", "Details": f"Outlook draft created ({os.path.basename(matched_pdf)})"})

        # 2. Fallback to generating .eml file if Outlook COM failed or unavailable
        if not draft_created:
            ok_eml, eml_path = create_eml_file(to_email, cc_list, formatted_subject, formatted_body, matched_pdf, eml_output_dir, f"{insured_name}_{policy_no}")
            if ok_eml:
                success_count += 1
                logs.append({"Row": idx + 1, "Insured Name": insured_name, "Status": "Success (.EML File)", "Details": f"Saved file: {os.path.basename(eml_path)}"})
            else:
                failed_count += 1
                logs.append({"Row": idx + 1, "Insured Name": insured_name, "Status": "Failed", "Details": f"EML creation error: {eml_path}"})

        progress_bar.progress((idx + 1) / total_rows)

    # --- Summary Metrics Display ---
    st.markdown("### 📊 Summary Totals")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total Rows Processed", total_rows)
    m2.metric("Drafts Created", success_count, delta="Success")
    m3.metric("Skipped Rows", skipped_count, delta_color="off")
    m4.metric("Failed", failed_count, delta_color="inverse")

    st.markdown("---")
    st.subheader("Detailed Execution Log")
    st.dataframe(pd.DataFrame(logs), use_container_width=True)